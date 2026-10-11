import json
import os
import re
import base64
import hashlib
import hmac
import logging
import secrets
import time
import unicodedata
from io import BytesIO

import psycopg
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile, Response
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from database import (
    atualizar_aluno,
    criar_aluno,
    listar_analises,
    normalizar_semestre_ano_ingresso,
    obter_analise,
    obter_conteudo_historico,
    obter_extracao_historico,
    obter_conexao,
    salvar_extracao_historico,
    salvar_revisao_disciplinas,
)
from services.comparacao_carga_horaria import comparar_carga_horaria
from services.documentos import extrair_texto_documento, validar_historico, validar_plano_origem
from services.extracao_historico import extrair_disciplinas
from services.equivalencia_agent import get_agent


app = FastAPI(title="equivalencIA API", version="0.1.0")
logger = logging.getLogger(__name__)
allowed_origins = os.getenv(
    "FRONTEND_ORIGINS",
    "http://localhost:5173,http://localhost:4173",
).split(",")
app.add_middleware(
    CORSMiddleware,
    allow_origins=[origin.strip() for origin in allowed_origins if origin.strip()],
    allow_credentials=True,
    allow_methods=["GET", "POST", "PUT", "PATCH", "OPTIONS"],
    allow_headers=["*"],
)

SESSION_COOKIE = "equivalencia_session"
SESSION_SECONDS = 8 * 60 * 60


def _verify_password(password: str, encoded: str) -> bool:
    try:
        algorithm, iterations, salt, expected = encoded.split("$")
        if algorithm != "pbkdf2_sha256":
            return False
        actual = hashlib.pbkdf2_hmac("sha256", password.encode(), bytes.fromhex(salt), int(iterations))
        return hmac.compare_digest(actual.hex(), expected)
    except (ValueError, TypeError):
        return False


def _session_secret() -> bytes:
    secret = os.getenv("SESSION_SECRET", "")
    if len(secret) < 32:
        raise HTTPException(status_code=503, detail="Configure SESSION_SECRET com pelo menos 32 caracteres.")
    return secret.encode()


def _make_session(user_id: int) -> str:
    payload = base64.urlsafe_b64encode(json.dumps({"sub": user_id, "exp": int(time.time()) + SESSION_SECONDS}).encode()).decode().rstrip("=")
    signature = hmac.new(_session_secret(), payload.encode(), hashlib.sha256).digest()
    return payload + "." + base64.urlsafe_b64encode(signature).decode().rstrip("=")


def _session_user(token: str | None):
    if not token or "." not in token:
        return None
    try:
        payload, signature = token.split(".", 1)
        expected = base64.urlsafe_b64encode(hmac.new(_session_secret(), payload.encode(), hashlib.sha256).digest()).decode().rstrip("=")
        if not hmac.compare_digest(signature, expected):
            return None
        decoded = json.loads(base64.urlsafe_b64decode(payload + "=" * (-len(payload) % 4)))
        if decoded.get("exp", 0) < time.time():
            return None
        with obter_conexao() as connection, connection.cursor() as cursor:
            cursor.execute(
                """SELECT c.id, c.nome, c.email, c.perfil, u.id, u.nome
                     FROM coordenador c JOIN unidade u ON u.id = c.unidade_id
                    WHERE c.id = %s AND c.ativo = TRUE AND u.ativa = TRUE""",
                (decoded["sub"],),
            )
            row = cursor.fetchone()
        return None if row is None else {"id": row[0], "name": row[1], "email": row[2], "role": row[3], "unit_id": row[4], "unit_name": row[5]}
    except Exception:
        return None


@app.middleware("http")
async def require_coordinator(request: Request, call_next):
    path = request.url.path
    public_paths = {"/api/health", "/api/auth/units", "/api/auth/login", "/api/auth/logout"}
    if path in public_paths or request.method == "OPTIONS":
        return await call_next(request)
    if path.startswith("/api/"):
        user = _session_user(request.cookies.get(SESSION_COOKIE))
        if user is None:
            from starlette.responses import JSONResponse
            return JSONResponse({"detail": "Acesso restrito à coordenação. Entre novamente."}, status_code=401)
        request.state.coordinator = user
        analysis_match = re.match(r"^/api/analyses/(\d+)(?:/|$)", path)
        if analysis_match:
            with obter_conexao() as connection, connection.cursor() as cursor:
                cursor.execute("SELECT unidade_id FROM analise WHERE id = %s", (int(analysis_match.group(1)),))
                owner = cursor.fetchone()
            if owner is None or owner[0] != user["unit_id"]:
                from starlette.responses import JSONResponse
                return JSONResponse({"detail": "Análise não encontrada."}, status_code=404)
    return await call_next(request)


class LoginInput(BaseModel):
    email: str
    password: str
    unit_id: int


@app.get("/api/auth/units")
def auth_units():
    with obter_conexao() as connection, connection.cursor() as cursor:
        cursor.execute("SELECT id, nome FROM unidade WHERE ativa = TRUE ORDER BY nome")
        return [{"id": row[0], "name": row[1]} for row in cursor.fetchall()]


@app.post("/api/auth/login")
def login(payload: LoginInput, response: Response):
    with obter_conexao() as connection, connection.cursor() as cursor:
        cursor.execute(
            """SELECT id, senha_hash FROM coordenador
                 WHERE lower(email) = lower(%s) AND unidade_id = %s AND ativo = TRUE""",
            (payload.email.strip(), payload.unit_id),
        )
        row = cursor.fetchone()
    if row is None or not _verify_password(payload.password, row[1]):
        raise HTTPException(status_code=401, detail="E-mail, senha ou unidade inválidos.")
    response.set_cookie(SESSION_COOKIE, _make_session(row[0]), max_age=SESSION_SECONDS, httponly=True, secure=os.getenv("SESSION_COOKIE_SECURE", "false").lower() == "true", samesite="strict", path="/")
    return {"ok": True}


@app.get("/api/auth/me")
def auth_me(request: Request):
    return request.state.coordinator


@app.post("/api/auth/logout")
def logout(response: Response):
    response.delete_cookie(SESSION_COOKIE, path="/")
    return {"ok": True}


class StudentInput(BaseModel):
    name: str = Field(min_length=1, max_length=150)
    ra: str | None = Field(default=None, max_length=30)


class HistoryReview(BaseModel):
    rows: list[dict]
    confirm: bool = False


class SuggestionDecision(BaseModel):
    decision: str
    note: str = ""


class ReportInput(BaseModel):
    coordinator_note: str = ""


class ManualSuggestionInput(BaseModel):
    source_id: int
    target_id: int
    rationale: str = Field(min_length=1, max_length=2000)


def _analysis_or_404(analysis_id: int):
    data = obter_analise(analysis_id)
    if data is None:
        raise HTTPException(status_code=404, detail="Análise não encontrada.")
    return data


def _suggestions(analysis_id: int):
    with obter_conexao() as connection, connection.cursor() as cursor:
        cursor.execute(
            """SELECT s.id, s.discipline_extra_id, s.target_discipline_id, s.confidence,
                      s.rationale, s.evidence, s.decision, s.coordinator_note, s.provider,
                      s.model, s.created_at, src.nome, dst.nome
                 FROM equivalencia_sugestao s
                 JOIN disciplina_extraida src ON src.id = s.discipline_extra_id
                 JOIN disciplina dst ON dst.id = s.target_discipline_id
                WHERE s.analysis_id = %s
                ORDER BY s.id""",
            (analysis_id,),
        )
        rows = cursor.fetchall()
    return [
        {
            "id": row[0],
            "source_id": row[1],
            "target_id": row[2],
            "confidence": float(row[3]) if row[3] is not None else None,
            "rationale": row[4],
            "evidence": row[5] or {},
            "decision": row[6],
            "note": row[7] or "",
            "provider": row[8],
            "model": row[9],
            "created_at": row[10].isoformat() if row[10] else None,
            "origin_name": row[11],
            "target_name": row[12],
            "evidence_summary": "; ".join(
                str(value) for value in (row[5] or {}).values() if value
            ),
        }
        for row in rows
    ]


def _analysis_payload(analysis_id: int):
    data = _analysis_or_404(analysis_id)
    extraction, disciplines = obter_extracao_historico(analysis_id)
    history = None
    plans = []
    with obter_conexao() as connection, connection.cursor() as cursor:
        cursor.execute(
            """SELECT id, tipo, nome_arquivo FROM documento
                 WHERE analise_id = %s ORDER BY id""",
            (analysis_id,),
        )
        documents = cursor.fetchall()
        history = next((item[2] for item in documents if item[1] == "historico"), None)
        plans = [item[2] for item in documents if item[1] == "plano_ensino_origem"]
        cursor.execute("SELECT nome FROM curso WHERE id = %s", (data[7],))
        target_course = cursor.fetchone()
        cursor.execute("SELECT codigo FROM matriz WHERE id = %s", (data[8],))
        target_matrix = cursor.fetchone()
    return {
        "id": data[0],
        "student_name": data[1],
        "student_ra": data[2],
        "term": data[3],
        "origin_status": data[4],
        "origin_institution": data[5],
        "created_at": data[6].isoformat() if data[6] else None,
        "target_course_id": data[7],
        "target_matrix_id": data[8],
        "student_id": data[9],
        "origin_course": data[10],
        "target_course": target_course[0] if target_course else None,
        "target_matrix": target_matrix[0] if target_matrix else None,
        "history_name": history,
        "plans": plans,
        "extraction": None if extraction is None else {
            "confirmed": extraction[4] is not None,
            "warning": extraction[3],
            "processed_at": extraction[2].isoformat() if extraction[2] else None,
            "rows": [
                {
                    "id": row[0], "code": row[1], "name": row[2], "grade": row[3],
                    "hours": row[4], "period": row[5], "status": row[6],
                    "evidence": row[7], "reviewed": row[8],
                }
                for row in disciplines
            ],
        },
        "suggestions": _suggestions(analysis_id),
    }


@app.get("/api/health")
def health():
    try:
        with obter_conexao() as connection, connection.cursor() as cursor:
            cursor.execute("SELECT 1")
            cursor.fetchone()
        return {"status": "ok", "database": "connected", "ai_configured": get_agent() is not None}
    except psycopg.Error as error:
        raise HTTPException(status_code=503, detail="Banco de dados indisponível.") from error


@app.get("/api/students")
def students(request: Request, query: str = "", limit: int = 30):
    limit = max(1, min(limit, 100))
    with obter_conexao() as connection, connection.cursor() as cursor:
        if query.strip():
            normalized = "".join(
                char for char in unicodedata.normalize("NFKD", query.strip().lower())
                if not unicodedata.combining(char)
            ).replace("!", "!!").replace("%", "!%").replace("_", "!_")
            search = "%" + normalized + "%"
            cursor.execute(
                """SELECT id, nome, ra FROM aluno
                    WHERE unidade_id = %s AND (translate(lower(nome), 'áàãâäéèêëíìîïóòõôöúùûüç', 'aaaaaeeeeiiiiooooouuuuc') LIKE %s ESCAPE '!'
                       OR COALESCE(ra, '') ILIKE %s ESCAPE '!')
                    ORDER BY nome LIMIT %s""",
                (request.state.coordinator["unit_id"], search, search, limit),
            )
        else:
            cursor.execute(
                "SELECT id, nome, ra FROM aluno WHERE unidade_id = %s ORDER BY nome LIMIT %s",
                (request.state.coordinator["unit_id"], limit),
            )
        rows = cursor.fetchall()
    return [{"id": row[0], "name": row[1], "ra": row[2]} for row in rows]


@app.post("/api/students", status_code=201)
def create_student(payload: StudentInput, request: Request):
    try:
        ident = criar_aluno(payload.name, payload.ra, request.state.coordinator["unit_id"])
        with obter_conexao() as connection, connection.cursor() as cursor:
            cursor.execute("SELECT id, nome, ra FROM aluno WHERE id = %s", (ident,))
            row = cursor.fetchone()
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    return {"id": row[0], "name": row[1], "ra": row[2]}


@app.put("/api/students/{student_id}")
def update_student(student_id: int, payload: StudentInput, request: Request):
    if not payload.name.strip():
        raise HTTPException(status_code=422, detail="O nome do aluno é obrigatório.")
    try:
        with obter_conexao() as connection, connection.cursor() as cursor:
            cursor.execute("UPDATE aluno SET nome = %s, ra = %s WHERE id = %s AND unidade_id = %s RETURNING id", (payload.name.strip(), payload.ra.strip() or None if payload.ra else None, student_id, request.state.coordinator["unit_id"]))
            updated = cursor.fetchone() is not None
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    if not updated:
        raise HTTPException(status_code=404, detail="Aluno cadastrado não encontrado.")
    with obter_conexao() as connection, connection.cursor() as cursor:
        cursor.execute("SELECT id, nome, ra FROM aluno WHERE id = %s AND unidade_id = %s", (student_id, request.state.coordinator["unit_id"]))
        row = cursor.fetchone()
    return {"id": row[0], "name": row[1], "ra": row[2]}


@app.get("/api/courses")
def courses(request: Request):
    with obter_conexao() as connection, connection.cursor() as cursor:
        cursor.execute("SELECT id, codigo, nome FROM curso WHERE unidade_id = %s ORDER BY nome", (request.state.coordinator["unit_id"],))
        rows = cursor.fetchall()
    return [{"id": row[0], "code": row[1], "name": row[2]} for row in rows]


@app.get("/api/courses/{course_id}/matrices")
def matrices(course_id: int, request: Request):
    with obter_conexao() as connection, connection.cursor() as cursor:
        cursor.execute(
            "SELECT m.id, m.codigo FROM matriz m JOIN curso c ON c.id = m.curso_id WHERE m.curso_id = %s AND c.unidade_id = %s ORDER BY m.codigo",
            (course_id, request.state.coordinator["unit_id"]),
        )
        rows = cursor.fetchall()
    return [{"id": row[0], "code": row[1]} for row in rows]


@app.get("/api/analyses")
def analyses(request: Request):
    with obter_conexao() as connection, connection.cursor() as cursor:
        cursor.execute("""SELECT a.id, al.nome, al.ra, a.semestre_ano, a.situacao_origem, a.procedencia,
                                 a.data_criacao, a.curso_id, a.matriz_id, a.aluno_id, a.curso_origem
                            FROM analise a JOIN aluno al ON al.id = a.aluno_id
                           WHERE a.unidade_id = %s ORDER BY a.data_criacao DESC, a.id DESC""", (request.state.coordinator["unit_id"],))
        records = cursor.fetchall()
    result = []
    for row in records:
        analysis_id = row[0]
        extraction, _ = obter_extracao_historico(analysis_id)
        suggestions = _suggestions(analysis_id)
        status = "review" if any(s["decision"] == "pending" for s in suggestions) else (
            "complete" if suggestions and all(s["decision"] in ("approved", "rejected") for s in suggestions)
            else "progress"
        )
        stage = "Aguardando revisão" if status == "review" else (
            "Concluída" if status == "complete" else "Em andamento"
        )
        with obter_conexao() as connection, connection.cursor() as cursor:
            cursor.execute("SELECT nome FROM curso WHERE id = %s", (row[7],))
            destination = cursor.fetchone()
        result.append({
            "id": analysis_id,
            "student_name": row[1],
            "origin_course": row[10],
            "target_course": destination[0] if destination else None,
            "status": status,
            "stage": stage,
            "updated_at": row[6].strftime("%d/%m/%Y") if row[6] else None,
        })
    return result


@app.post("/api/analyses", status_code=201)
async def create_analysis(
    request: Request,
    student_id: int = Form(),
    origin_institution: str = Form(),
    origin_course: str = Form(),
    origin_status: str = Form(),
    target_course_id: int = Form(),
    target_matrix_id: int = Form(),
    term: str = Form(""),
    history: UploadFile = File(),
    plans: list[UploadFile] | None = File(default=None),
):
    max_bytes = int(os.getenv("MAX_DOCUMENT_BYTES", str(20 * 1024 * 1024)))
    history_bytes = await history.read(max_bytes + 1)
    if len(history_bytes) > max_bytes:
        raise HTTPException(status_code=413, detail="O histórico excede o limite de 20 MB.")
    try:
        validar_historico(history.filename or "", history_bytes)
        normalized_plans = []
        for file in plans or []:
            content = await file.read(max_bytes + 1)
            if len(content) > max_bytes:
                raise HTTPException(
                    status_code=413,
                    detail="O plano " + (file.filename or "") + " excede o limite de 20 MB.",
                )
            validar_plano_origem(file.filename or "", content)
            normalized_plans.append((file.filename or "plano", content))
        term = normalizar_semestre_ano_ingresso(term)
        with obter_conexao() as connection, connection.cursor() as cursor:
            unit_id = request.state.coordinator["unit_id"]
            cursor.execute("SELECT 1 FROM aluno WHERE id = %s AND unidade_id = %s", (student_id, unit_id))
            if cursor.fetchone() is None:
                raise HTTPException(status_code=404, detail="Aluno cadastrado não encontrado.")
            cursor.execute(
                "SELECT 1 FROM matriz m JOIN curso c ON c.id = m.curso_id WHERE m.id = %s AND m.curso_id = %s AND c.unidade_id = %s",
                (target_matrix_id, target_course_id, unit_id),
            )
            if cursor.fetchone() is None:
                raise HTTPException(status_code=422, detail="A matriz selecionada não pertence ao curso de destino.")
            cursor.execute(
                """INSERT INTO analise
                     (unidade_id, aluno_id, semestre_ano, curso_origem, situacao_origem,
                      procedencia, curso_id, matriz_id)
                   VALUES (%s, %s, %s, %s, %s, %s, %s, %s) RETURNING id""",
                (
                    unit_id, student_id, term, origin_course.strip() or None,
                    origin_status.strip() or None, origin_institution.strip() or None,
                    target_course_id, target_matrix_id,
                ),
            )
            analysis_id = cursor.fetchone()[0]
            cursor.execute(
                """INSERT INTO documento (analise_id, tipo, nome_arquivo, conteudo)
                   VALUES (%s, 'historico', %s, %s)""",
                (analysis_id, history.filename or "historico", history_bytes),
            )
            if normalized_plans:
                cursor.executemany(
                    """INSERT INTO documento (analise_id, tipo, nome_arquivo, conteudo)
                       VALUES (%s, 'plano_ensino_origem', %s, %s)""",
                    [(analysis_id, filename, content) for filename, content in normalized_plans],
                )
        payload = _analysis_payload(analysis_id)
        payload["analysis_id"] = analysis_id
        return payload
    except HTTPException:
        raise
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    except psycopg.Error as error:
        logger.exception("Falha de banco ao salvar análise de aproveitamento.")
        raise HTTPException(status_code=503, detail="Não foi possível salvar a análise.") from error


@app.get("/api/analyses/{analysis_id}")
def read_analysis(analysis_id: int):
        result = _analysis_payload(analysis_id)
        result["analysis_id"] = analysis_id
        return result


@app.post("/api/analyses/{analysis_id}/extract")
def extract_history(analysis_id: int):
    _analysis_or_404(analysis_id)
    document = obter_conteudo_historico(analysis_id)
    if document is None:
        raise HTTPException(status_code=409, detail="Anexe o histórico acadêmico antes da extração.")
    try:
        disciplines, warning = extrair_disciplinas(document[1], document[2])
        salvar_extracao_historico(analysis_id, document[0], disciplines, warning)
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    return _analysis_payload(analysis_id)


@app.put("/api/analyses/{analysis_id}/history-review")
def review_history(analysis_id: int, payload: HistoryReview):
    _analysis_or_404(analysis_id)
    rows = [
        {
            "ID": row.get("id"),
            "Código": row.get("code"),
            "Disciplina": row.get("name"),
            "Nota": row.get("grade"),
            "Carga horária": row.get("hours"),
            "Ano/semestre": row.get("period"),
            "Situação": row.get("status"),
        }
        for row in payload.rows
    ]
    try:
        saved = salvar_revisao_disciplinas(analysis_id, rows, confirmar=payload.confirm)
        if not saved:
            raise HTTPException(status_code=409, detail="Execute a extração antes de revisar os dados.")
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    return _analysis_payload(analysis_id)


@app.post("/api/analyses/{analysis_id}/ai-analysis")
def run_ai_analysis(analysis_id: int):
    data = _analysis_payload(analysis_id)
    if not data["extraction"] or not data["extraction"]["confirmed"]:
        raise HTTPException(status_code=409, detail="O coordenador deve revisar e confirmar o histórico antes da análise.")
    agent = get_agent()
    if agent is None:
        return {
            "suggestions": [],
            "ai_configured": False,
            "message": "Agente de IA ainda não conectado. Você pode adicionar correspondências manualmente e concluir a revisão.",
        }
    context = _build_agent_context(analysis_id)
    try:
        suggestions = agent.analyze(context)
        _validate_agent_suggestions(suggestions, context)
        _save_suggestions(analysis_id, suggestions, agent.provider_name, agent.model_name)
    except Exception as error:
        raise HTTPException(status_code=502, detail="O agente de IA não conseguiu concluir a análise.") from error
    return {"suggestions": _suggestions(analysis_id)}


@app.get("/api/analyses/{analysis_id}/candidates")
def list_candidates(analysis_id: int):
    data = _analysis_payload(analysis_id)
    if not data["extraction"] or not data["extraction"]["confirmed"]:
        raise HTTPException(status_code=409, detail="Confirme os dados extraídos antes de escolher disciplinas.")
    return _build_agent_context(analysis_id)["candidates"]


@app.post("/api/analyses/{analysis_id}/manual-suggestions", status_code=201)
def add_manual_suggestion(analysis_id: int, payload: ManualSuggestionInput):
    data = _analysis_payload(analysis_id)
    if not data["extraction"] or not data["extraction"]["confirmed"]:
        raise HTTPException(status_code=409, detail="Confirme os dados extraídos antes de adicionar equivalências.")
    if not payload.rationale.strip():
        raise HTTPException(status_code=422, detail="Informe a justificativa da correspondência manual.")
    candidate_pairs = {
        (item["source"]["id"], item["target"]["id"])
        for item in _build_agent_context(analysis_id)["candidates"]
    }
    if (payload.source_id, payload.target_id) not in candidate_pairs:
        raise HTTPException(
            status_code=422,
            detail="A correspondência não atende aos critérios de aprovação e carga horária.",
        )
    with obter_conexao() as connection, connection.cursor() as cursor:
        cursor.execute(
            """INSERT INTO equivalencia_sugestao
                 (analysis_id, discipline_extra_id, target_discipline_id,
                  rationale, evidence, decision, provider, model)
               VALUES (%s, %s, %s, %s, %s::jsonb, 'pending', 'manual', NULL)
               ON CONFLICT (analysis_id, discipline_extra_id, target_discipline_id)
               DO UPDATE SET rationale = EXCLUDED.rationale,
                             evidence = EXCLUDED.evidence,
                             decision = 'pending', coordinator_note = '',
                             provider = 'manual', model = NULL,
                             created_at = CURRENT_TIMESTAMP, reviewed_at = NULL""",
            (
                analysis_id, payload.source_id, payload.target_id,
                payload.rationale.strip(),
                json.dumps({"origin": "coordinator_manual_entry"}),
            ),
        )
    return _analysis_payload(analysis_id)


def _build_agent_context(analysis_id: int):
    data = _analysis_or_404(analysis_id)
    extraction, extracted = obter_extracao_historico(analysis_id)
    eligible = [
        row for row in extracted
        if row[8] and (row[6] or "").strip().lower().startswith("aprovad")
    ]
    with obter_conexao() as connection, connection.cursor() as cursor:
        cursor.execute(
            """SELECT d.id, d.codigo, d.nome, d.carga_horaria, d.periodo,
                      p.ementa, p.conteudo_programatico, p.objetivos
                 FROM disciplina d
                 LEFT JOIN plano_ensino p ON p.disciplina_id = d.id
                WHERE d.matriz_id = %s
                ORDER BY d.periodo, d.nome""",
            (data[8],),
        )
        targets = cursor.fetchall()
        cursor.execute(
            """SELECT nome_arquivo, conteudo FROM documento
                WHERE analise_id = %s AND tipo = 'plano_ensino_origem'
                ORDER BY id""",
            (analysis_id,),
        )
        plans = cursor.fetchall()
    candidates = []
    for source in eligible:
        for target in targets:
            hours = comparar_carga_horaria(source[4], target[3])
            if hours["atendido"] is not True:
                continue
            candidates.append({
                "source": {
                    "id": source[0], "code": source[1], "name": source[2],
                    "grade": source[3], "hours": source[4], "term": source[5],
                    "status": source[6], "evidence": source[7],
                },
                "target": {
                    "id": target[0], "code": target[1], "name": target[2],
                    "hours": target[3], "term": target[4], "syllabus": target[5],
                    "content": target[6], "objectives": target[7],
                },
                "workload_check": {
                    "eligible": hours["atendido"],
                    "source_hours": str(hours["origem"]) if hours["origem"] is not None else None,
                    "target_hours": str(hours["destino"]) if hours["destino"] is not None else None,
                    "reason": hours["motivo"],
                },
            })
    return {
        "analysis_id": analysis_id,
        "student": {"id": data[9], "name": data[1], "ra": data[2]},
        "origin": {"institution": data[5], "course": data[10]},
        "destination": {"course_id": data[7], "matrix_id": data[8]},
        "source_plans": [
            {"name": name, "text": extrair_texto_documento(name, content)}
            for name, content in plans
        ],
        "candidates": candidates,
        "constraints": [
            "Only confirmed and explicitly approved source disciplines are eligible.",
            "Source workload must be equal to or greater than target workload.",
            "Never invent data; cite evidence from the source documents and target syllabus.",
            "Return suggestions only; the coordinator makes the final decision.",
        ],
    }


def _save_suggestions(analysis_id, suggestions, provider, model):
    with obter_conexao() as connection, connection.cursor() as cursor:
        cursor.execute("DELETE FROM equivalencia_sugestao WHERE analysis_id = %s", (analysis_id,))
        for suggestion in suggestions:
            cursor.execute(
                """INSERT INTO equivalencia_sugestao
                     (analysis_id, discipline_extra_id, target_discipline_id,
                      confidence, rationale, evidence, decision, provider, model)
                   VALUES (%s, %s, %s, %s, %s, %s::jsonb, 'pending', %s, %s)
                   ON CONFLICT (analysis_id, discipline_extra_id, target_discipline_id)
                   DO UPDATE SET confidence = EXCLUDED.confidence,
                                 rationale = EXCLUDED.rationale,
                                 evidence = EXCLUDED.evidence,
                                 decision = 'pending',
                                 coordinator_note = '',
                                 provider = EXCLUDED.provider,
                                 model = EXCLUDED.model,
                                 created_at = CURRENT_TIMESTAMP,
                                 reviewed_at = NULL""",
                (
                    analysis_id,
                    suggestion["source_id"],
                    suggestion["target_id"],
                    suggestion.get("confidence"),
                    suggestion["rationale"],
                    json.dumps(suggestion.get("evidence", {}), ensure_ascii=False),
                    provider,
                    model,
                ),
            )


def _validate_agent_suggestions(suggestions, context):
    if not isinstance(suggestions, list):
        raise ValueError("O agente deve retornar uma lista estruturada de sugestões.")
    allowed = {
        (candidate["source"]["id"], candidate["target"]["id"])
        for candidate in context["candidates"]
    }
    seen = set()
    for suggestion in suggestions:
        pair = (suggestion.get("source_id"), suggestion.get("target_id"))
        if pair not in allowed or pair in seen:
            raise ValueError("O agente retornou uma relação fora das evidências elegíveis.")
        confidence = suggestion.get("confidence")
        if confidence is not None and not 0 <= float(confidence) <= 1:
            raise ValueError("Confiança deve estar entre 0 e 1.")
        if not str(suggestion.get("rationale", "")).strip():
            raise ValueError("Cada sugestão precisa explicar seu motivo.")
        if not isinstance(suggestion.get("evidence", {}), dict):
            raise ValueError("As evidências devem ser retornadas em formato estruturado.")
        seen.add(pair)


@app.patch("/api/analyses/{analysis_id}/suggestions/{suggestion_id}")
def decide_suggestion(analysis_id: int, suggestion_id: int, payload: SuggestionDecision):
    _analysis_or_404(analysis_id)
    if payload.decision not in {"approved", "rejected", "edited"}:
        raise HTTPException(status_code=422, detail="Decisão inválida.")
    with obter_conexao() as connection, connection.cursor() as cursor:
        cursor.execute(
            """UPDATE equivalencia_sugestao
                  SET decision = %s, coordinator_note = %s, reviewed_at = CURRENT_TIMESTAMP
                WHERE id = %s AND analysis_id = %s""",
            (payload.decision, payload.note, suggestion_id, analysis_id),
        )
        if cursor.rowcount != 1:
            raise HTTPException(status_code=404, detail="Sugestão não encontrada.")
    return _analysis_payload(analysis_id)


@app.post("/api/analyses/{analysis_id}/report")
def report(analysis_id: int, payload: ReportInput):
    data = _analysis_payload(analysis_id)
    approved = [item for item in data["suggestions"] if item["decision"] == "approved"]
    try:
        from docx import Document
        template_path = os.getenv("REPORT_TEMPLATE_PATH", "documents/modelo_parecer.docx")
        document = Document(template_path) if os.path.isfile(template_path) else Document()
        placeholders = {
            "{{student_name}}": data["student_name"] or "Não informado",
            "{{student_ra}}": data["student_ra"] or "Não informado",
            "{{origin_institution}}": data["origin_institution"] or "Não informado",
            "{{origin_course}}": data["origin_course"] or "Não informado",
            "{{target_course}}": data["target_course"] or "Não informado",
            "{{target_matrix}}": data["target_matrix"] or "Não informado",
            "{{term}}": data["term"] or "Não informado",
            "{{coordinator_note}}": payload.coordinator_note or "",
        }
        for paragraph in document.paragraphs:
            for token, value in placeholders.items():
                if token in paragraph.text:
                    paragraph.text = paragraph.text.replace(token, value)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for token, value in placeholders.items():
                        if token in cell.text:
                            cell.text = cell.text.replace(token, value)
        document.add_heading("Parecer de aproveitamento de estudos", 0)
        document.add_paragraph("Documento gerado para revisão da coordenação acadêmica.")
        document.add_heading("Identificação do aluno", level=1)
        document.add_paragraph("Nome: " + data["student_name"])
        document.add_paragraph("RA: " + (data["student_ra"] or "Não informado"))
        document.add_paragraph("Instituição de origem: " + (data["origin_institution"] or "Não informado"))
        document.add_paragraph("Curso de origem: " + (data["origin_course"] or "Não informado"))
        document.add_paragraph("Curso de destino: " + (data["target_course"] or "Não informado"))
        document.add_paragraph("Matriz: " + (data["target_matrix"] or "Não informado"))
        document.add_heading("Resultado da análise", level=1)
        if data["suggestions"]:
            table = document.add_table(rows=1, cols=4)
            table.style = "Table Grid"
            for cell, label in zip(table.rows[0].cells, ("Origem", "Destino", "Confiança", "Situação")):
                cell.text = label
            labels = {"approved": "Aprovada", "rejected": "Não aprovada", "pending": "Pendente"}
            for item in data["suggestions"]:
                row = table.add_row().cells
                row[0].text = item["origin_name"]
                row[1].text = item["target_name"]
                row[2].text = str(item["confidence"] or "")
                row[3].text = labels.get(item["decision"], "Pendente")
        else:
            document.add_paragraph("Nenhuma equivalência foi sugerida ou registrada manualmente.")
        if not approved:
            document.add_paragraph("Nenhuma disciplina foi dispensada nesta análise.")
        document.add_paragraph(payload.coordinator_note or "")
        document.add_paragraph("Este parecer deve ser conferido e assinado conforme os procedimentos institucionais.")
        output = BytesIO()
        document.save(output)
        output.seek(0)
    except Exception as error:
        raise HTTPException(status_code=500, detail="Não foi possível gerar o documento.") from error
    filename = "parecer-aproveitamento-" + str(analysis_id) + ".docx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=" + filename},
    )


# In the hosted build, serve the React application from the same origin as the
# API. This keeps session cookies simple and avoids separate CORS configuration.
frontend_dist = os.path.join(os.path.dirname(os.path.dirname(__file__)), "frontend", "dist")
if os.path.isdir(frontend_dist):
    app.mount("/", StaticFiles(directory=frontend_dist, html=True), name="frontend")
