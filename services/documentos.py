from io import BytesIO
from pathlib import PurePath
from zipfile import BadZipFile, ZipFile
from docx import Document
from pypdf import PdfReader


FORMATOS_HISTORICO = ("pdf", "docx")
FORMATOS_PLANO_ORIGEM = ("pdf", "docx")


def validar_historico(nome_arquivo, conteudo):
    extensao = PurePath(nome_arquivo or "").suffix.lower().lstrip(".")
    if extensao not in FORMATOS_HISTORICO:
        raise ValueError("Formato não permitido. Envie um arquivo PDF ou DOCX.")
    if not conteudo:
        raise ValueError("O arquivo está vazio. Selecione um histórico válido.")
    if extensao == "pdf" and not conteudo.startswith(b"%PDF-"):
        raise ValueError("O arquivo não parece ser um PDF válido.")
    if extensao == "docx":
        try:
            with ZipFile(BytesIO(conteudo)) as arquivo:
                nomes = set(arquivo.namelist())
        except BadZipFile:
            raise ValueError("O arquivo não parece ser um DOCX válido.") from None
        if "[Content_Types].xml" not in nomes or "word/document.xml" not in nomes:
            raise ValueError("O arquivo não parece ser um DOCX válido.")
    return extensao


def validar_plano_origem(nome_arquivo, conteudo):
    extensao = PurePath(nome_arquivo or "").suffix.lower().lstrip(".")
    if extensao not in FORMATOS_PLANO_ORIGEM:
        raise ValueError("Formato não permitido. Envie os planos em PDF ou DOCX.")
    if not conteudo:
        raise ValueError(f"O arquivo {nome_arquivo} está vazio.")
    if extensao == "pdf" and not conteudo.startswith(b"%PDF-"):
        raise ValueError(f"O arquivo {nome_arquivo} não parece ser um PDF válido.")
    if extensao == "docx":
        try:
            with ZipFile(BytesIO(conteudo)) as arquivo:
                nomes = set(arquivo.namelist())
        except BadZipFile:
            raise ValueError(f"O arquivo {nome_arquivo} não parece ser um DOCX válido.") from None
        if "[Content_Types].xml" not in nomes or "word/document.xml" not in nomes:
            raise ValueError(f"O arquivo {nome_arquivo} não parece ser um DOCX válido.")
    return extensao


def extrair_texto_documento(nome_arquivo, conteudo):
    """Read text from supported source-plan files without inventing missing content."""
    extensao = validar_plano_origem(nome_arquivo, conteudo)
    if extensao == "pdf":
        pdf = PdfReader(BytesIO(conteudo), strict=False)
        return "\n".join(page.extract_text() or "" for page in pdf.pages).strip()
    doc = Document(BytesIO(conteudo))
    paragraphs = [item.text for item in doc.paragraphs if item.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(paragraphs).strip()
